import streamlit as st
import PyPDF2
from docx import Document
from io import BytesIO
import base64
from fpdf import FPDF
from .schemas import ResumeAnalysis
from .ollama_handler import structured_ollama_call


def parse_resume(file) -> str:
    """Extract text from uploaded PDF or DOCX file."""
    if file.type == "application/pdf":
        pdf = PyPDF2.PdfReader(file)
        return "\n".join([p.extract_text() or "" for p in pdf.pages])
    elif file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        doc = Document(file)
        return "\n".join([p.text for p in doc.paragraphs])
    return ""


def analyze_resume(jd: str, resume_text: str) -> ResumeAnalysis:
    """Call the LLM to analyze resume against the job description and return structured data."""
    prompt = f"""
Analyze this resume against the job description.

Job Description:
{jd}

Resume:
{resume_text}

Return a JSON object with the following fields:
- name: candidate's full name
- contact_info: email or phone number
- experience_summary: a brief summary of total experience
- match_score: integer score out of 100
- is_good_fit: true if match_score >= 60, else false
- strengths: list of key strengths
- weaknesses: list of areas for improvement
- missing_keywords: list of important keywords missing
- score_breakdown: dict with sub-scores
- detailed_report: full analysis narrative
"""
    return structured_ollama_call(
        prompt=prompt,
        response_model=ResumeAnalysis,
        model="gemma3"
    )


def generate_pdf(report: ResumeAnalysis) -> bytes:
    """Generate a PDF report from the ResumeAnalysis data, sanitizing Unicode to Latin-1, returning bytes."""
    def safe(text: str) -> str:
        # replace characters not encodable in latin-1
        return text.encode('latin-1', 'replace').decode('latin-1')

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Arial", 'B', 16)
    pdf.cell(0, 10, safe(f"Candidate: {report.name}"), ln=True)
    pdf.set_font("Arial", size=12)
    pdf.cell(0, 8, safe(f"Contact: {report.contact_info}"), ln=True)
    pdf.cell(0, 8, safe(f"Experience: {report.experience_summary}"), ln=True)
    pdf.cell(0, 8, safe(f"Match Score: {report.match_score}/100"), ln=True)
    verdict = safe("Good Fit" if report.is_good_fit else "Not a Good Fit")
    pdf.cell(0, 8, safe(f"Verdict: {verdict}"), ln=True)
    pdf.ln(5)
    pdf.set_font("Arial", 'B', 14)
    pdf.cell(0, 10, safe("Strengths:"), ln=True)
    pdf.set_font("Arial", size=12)
    for s in report.strengths:
        pdf.multi_cell(0, 8, safe(f"- {s}"))
    pdf.ln(3)
    pdf.set_font("Arial", 'B', 14)
    pdf.cell(0, 10, safe("Weaknesses:"), ln=True)
    pdf.set_font("Arial", size=12)
    for w in report.weaknesses:
        pdf.multi_cell(0, 8, safe(f"- {w}"))
    pdf.ln(3)
    pdf.set_font("Arial", 'B', 14)
    pdf.cell(0, 10, safe("Missing Keywords:"), ln=True)
    pdf.set_font("Arial", size=12)
    for kw in report.missing_keywords:
        pdf.multi_cell(0, 8, safe(f"- {kw}"))
    pdf.ln(5)
    pdf.set_font("Arial", 'B', 14)
    pdf.cell(0, 10, safe("Detailed Analysis:"), ln=True)
    pdf.set_font("Arial", size=12)
    for line in report.detailed_report.split("\n"):
        pdf.multi_cell(0, 8, safe(line))
    # output returns a string, so encode to bytes
    pdf_str = pdf.output(dest="S")
    return pdf_str.encode('latin-1')


def show_ui():
    st.header("Resume Analyzer 📄")
    jd = st.text_area("Job Description", height=150)
    resumes = st.file_uploader(
        "Upload Resumes", type=["pdf", "docx"], accept_multiple_files=True
    )

    if "resume_results" not in st.session_state:
        st.session_state.resume_results = {}

    if st.button("Analyze") and jd and resumes:
        for resume in resumes:
            if resume.name not in st.session_state.resume_results:
                with st.spinner(f"Analyzing {resume.name}…"):
                    text = parse_resume(resume)
                    st.session_state.resume_results[resume.name] = analyze_resume(jd, text)

    for name, result in st.session_state.resume_results.items():
        with st.expander(name, expanded=True):
            # Display key candidate info
            st.subheader(result.name)
            st.write(f"📞 Contact: {result.contact_info}")
            st.write(f"🧑‍💼 Experience: {result.experience_summary}")
            fit_label = "✅ Good Fit" if result.is_good_fit else "❌ Not a Good Fit"
            st.metric("Overall Fit", fit_label)
            st.write(f"**Match Score:** {result.match_score}/100")

            # Download detailed PDF report
            pdf_bytes = generate_pdf(result)
            b64 = base64.b64encode(pdf_bytes).decode()
            href = f'<a href="data:application/octet-stream;base64,{b64}" download="{name}_analysis.pdf">📥 Download Detailed Report (PDF)</a>'
            st.markdown(href, unsafe_allow_html=True)
            st.markdown("---")
